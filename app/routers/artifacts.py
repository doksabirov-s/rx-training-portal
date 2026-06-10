import copy
import json
import tempfile
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, Form, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse
from notebooklm import NotebookLMClient, SlideDeckFormat

from app.dependencies import get_client

router = APIRouter()

_DOWNLOADS = Path(tempfile.gettempdir()) / "rx_portal_downloads"
_DOWNLOADS.mkdir(exist_ok=True)
_SLIDES_DB = _DOWNLOADS / "slides_history.json"

# In-memory generation status: notebook_id -> message string
_gen_status: dict[str, str] = {}


def _load_history() -> dict:
    if _SLIDES_DB.exists():
        try:
            return json.loads(_SLIDES_DB.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def _save_slide_record(notebook_id: str, path: Path, topic: str, parts: int) -> None:
    history = _load_history()
    records = history.get(notebook_id, [])
    records.append({
        "filename": path.name,
        "topic": topic.strip() or "Без темы",
        "parts": parts,
        "slides": parts * 16,
        "created_at": datetime.now().strftime("%d.%m.%Y %H:%M"),
    })
    history[notebook_id] = records
    _SLIDES_DB.write_text(json.dumps(history, ensure_ascii=False, indent=2), encoding="utf-8")


def _remove_logo(pptx_path: Path) -> None:
    """Remove NotebookLM logo watermark (picture shapes) from every slide."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(str(pptx_path))
    changed = False
    for slide in prs.slides:
        sp_tree = slide.shapes._spTree
        for pic in list(sp_tree.findall(qn("p:pic"))):
            sp_tree.remove(pic)
            changed = True
    if changed:
        prs.save(str(pptx_path))


def _rewrite_zip(zip_path: Path, updates: dict) -> None:
    """Rewrite a ZIP file, replacing or adding entries from updates dict."""
    import zipfile, os
    tmp = zip_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(zip_path, "r") as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item.filename, updates.pop(item.filename, None) or src.read(item.filename))
        for path, data in updates.items():
            dst.writestr(path, data)
    os.replace(str(tmp), str(zip_path))


def _merge_pptx(files: list[Path]) -> Path:
    """Merge PPTX files via direct ZIP manipulation (no python-pptx object model)."""
    import zipfile, shutil, re

    if len(files) == 1:
        return files[0]

    SLIDE_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
    CT_SLIDE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"

    out = _DOWNLOADS / f"{files[0].stem}_merged.pptx"
    shutil.copy(str(files[0]), str(out))

    for src_file in files[1:]:
        with zipfile.ZipFile(out, "r") as z:
            names = z.namelist()
            prs_rels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
            prs_xml  = z.read("ppt/presentation.xml").decode("utf-8")
            ct_xml   = z.read("[Content_Types].xml").decode("utf-8")

        base_count = len([n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)])
        max_rid = max((int(m) for m in re.findall(r'Id="rId(\d+)"', prs_rels)), default=0)
        max_sid = max((int(m) for m in re.findall(r'<p:sldId[^>]+id="(\d+)"', prs_xml)), default=255)

        updates: dict[str, bytes] = {}

        with zipfile.ZipFile(src_file, "r") as sz:
            src_names = sz.namelist()
            src_slides = sorted(
                [n for n in src_names if re.match(r"ppt/slides/slide\d+\.xml$", n)],
                key=lambda x: int(re.search(r"\d+", x).group()),
            )
            for i, slide_path in enumerate(src_slides):
                new_num  = base_count + i + 1
                new_slide = f"ppt/slides/slide{new_num}.xml"
                new_rels  = f"ppt/slides/_rels/slide{new_num}.xml.rels"

                updates[new_slide] = sz.read(slide_path)

                src_rels_name = slide_path.replace("slides/", "slides/_rels/") + ".rels"
                if src_rels_name in src_names:
                    updates[new_rels] = sz.read(src_rels_name)

                max_rid += 1; max_sid += 1
                rid = f"rId{max_rid}"

                prs_rels = prs_rels.replace(
                    "</Relationships>",
                    f'<Relationship Id="{rid}" Type="{SLIDE_TYPE}" '
                    f'Target="slides/slide{new_num}.xml"/></Relationships>',
                )
                prs_xml = prs_xml.replace(
                    "</p:sldIdLst>",
                    f'<p:sldId id="{max_sid}" r:id="{rid}"/></p:sldIdLst>',
                )
                ct_xml = ct_xml.replace(
                    "</Types>",
                    f'<Override PartName="/ppt/slides/slide{new_num}.xml" '
                    f'ContentType="{CT_SLIDE}"/></Types>',
                )

        updates["ppt/_rels/presentation.xml.rels"] = prs_rels.encode("utf-8")
        updates["ppt/presentation.xml"]            = prs_xml.encode("utf-8")
        updates["[Content_Types].xml"]             = ct_xml.encode("utf-8")
        _rewrite_zip(out, updates)

    return out


async def _research_and_import(client: NotebookLMClient, notebook_id: str, topic: str):
    """Search web for topic and import found sources into the notebook."""
    research = await client.research.start(notebook_id, topic, source="web", mode="deep")
    if not research:
        return
    task = await client.research.wait_for_completion(
        notebook_id, research.task_id, timeout=300.0
    )
    if task and task.sources:
        sources = [{"url": s.url, "title": s.title} for s in task.sources if s.url]
        if sources:
            await client.research.import_sources_with_verification(
                notebook_id, research.task_id, sources
            )


async def _run_extended_slides(
    app, notebook_id: str, parts: int, topic: str = "", auto_search: bool = False
):
    """Optionally search web, then generate `parts` slide decks and merge."""
    client: NotebookLMClient = app.state.notebooklm

    if auto_search and topic.strip():
        _gen_status[notebook_id] = f"Поиск источников в интернете по теме: «{topic[:60]}»…"
        try:
            await _research_and_import(client, notebook_id, topic)
        except Exception as exc:
            _gen_status[notebook_id] = f"Поиск завершён с ошибкой ({exc}), продолжаю генерацию…"

    topic_prefix = f"Тема: {topic}. " if topic.strip() else ""
    part_suffixes = [
        "Введение, ключевые определения и основные концепции. Язык презентации: русский.",
        "Детальный анализ, данные, механизмы и доказательная база. Язык презентации: русский.",
        "Клиническое применение, кейсы и практические рекомендации. Язык презентации: русский.",
        "Выводы, перспективы и дополнительная информация. Язык презентации: русский.",
    ]

    pptx_files: list[Path] = []
    try:
        for i in range(parts):
            _gen_status[notebook_id] = f"Extended slides: генерирую часть {i + 1} из {parts}…"
            instructions = topic_prefix + part_suffixes[i % len(part_suffixes)]
            task = await client.artifacts.generate_slide_deck(
                notebook_id,
                language="ru",
                instructions=instructions,
                slide_format=SlideDeckFormat.DETAILED_DECK,
            )
            await client.artifacts.wait_for_completion(notebook_id, task.task_id, timeout=600.0)

            out = _DOWNLOADS / f"{notebook_id}_slides_part{i + 1}.pptx"
            await client.artifacts.download_slide_deck(notebook_id, out, output_format="pptx")
            pptx_files.append(out)

        _gen_status[notebook_id] = "Extended slides: объединяю файлы…"
        merged = _merge_pptx(pptx_files)
        _save_slide_record(notebook_id, merged, topic, parts)
        _gen_status[notebook_id] = f"__slides_ready__{merged}"
    except Exception as exc:
        _gen_status[notebook_id] = f"Ошибка генерации: {exc}"


async def _run_generation(app, notebook_id: str, kind: str, **kwargs):
    client: NotebookLMClient = app.state.notebooklm
    _gen_status[notebook_id] = f"{kind} generation in progress…"
    try:
        if kind == "flashcards":
            task = await client.artifacts.generate_flashcards(notebook_id, **kwargs)
        elif kind == "study-guide":
            task = await client.artifacts.generate_study_guide(notebook_id)
        elif kind == "audio":
            task = await client.artifacts.generate_audio(notebook_id)
        else:
            _gen_status[notebook_id] = f"Unknown artifact kind: {kind}"
            return
        await client.artifacts.wait_for_completion(notebook_id, task.task_id, timeout=1200.0)
        _gen_status.pop(notebook_id, None)
    except Exception as exc:
        _gen_status[notebook_id] = f"Generation failed: {exc}"


@router.post("/notebooks/{notebook_id}/artifacts/flashcards")
async def generate_flashcards(
    notebook_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    count: int = Form(10),
    client: NotebookLMClient = Depends(get_client),
):
    background_tasks.add_task(_run_generation, request.app, notebook_id, "flashcards", count=count)
    return RedirectResponse(
        url=f"/notebooks/{notebook_id}?status=Flashcard+generation+started.+Refresh+to+check+status.",
        status_code=303,
    )


@router.post("/notebooks/{notebook_id}/artifacts/study-guide")
async def generate_study_guide(
    notebook_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    client: NotebookLMClient = Depends(get_client),
):
    background_tasks.add_task(_run_generation, request.app, notebook_id, "study-guide")
    return RedirectResponse(
        url=f"/notebooks/{notebook_id}?status=Study+guide+generation+started.+Refresh+to+check+status.",
        status_code=303,
    )


@router.post("/notebooks/{notebook_id}/artifacts/audio")
async def generate_audio(
    notebook_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    client: NotebookLMClient = Depends(get_client),
):
    background_tasks.add_task(_run_generation, request.app, notebook_id, "audio")
    return RedirectResponse(
        url=f"/notebooks/{notebook_id}?status=Podcast+generation+started+%28takes+up+to+20+min%29.+Refresh+to+check+status.",
        status_code=303,
    )


@router.post("/notebooks/{notebook_id}/artifacts/slides")
async def generate_extended_slides(
    notebook_id: str,
    request: Request,
    background_tasks: BackgroundTasks,
    parts: int = Form(2),
    topic: str = Form(""),
    auto_search: str = Form(""),
    client: NotebookLMClient = Depends(get_client),
):
    parts = max(1, min(parts, 4))  # clamp 1–4
    do_search = bool(auto_search) and bool(topic.strip())
    background_tasks.add_task(
        _run_extended_slides, request.app, notebook_id, parts, topic, do_search
    )
    slides = parts * 16
    return RedirectResponse(
        url=f"/notebooks/{notebook_id}?status=Generating+{slides}+slides+in+{parts}+parts.+Refresh+to+check+status.",
        status_code=303,
    )


@router.get("/notebooks/{notebook_id}/artifacts/slides/download")
async def download_extended_slides(notebook_id: str):
    status = _gen_status.get(notebook_id, "")
    if status.startswith("__slides_ready__"):
        path = Path(status.removeprefix("__slides_ready__"))
        if path.exists():
            return FileResponse(
                path=str(path),
                media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                filename="presentation.pptx",
            )
    raise HTTPException(status_code=404, detail="Slides not ready yet. Please wait and refresh.")


@router.get("/notebooks/{notebook_id}/generation-status")
async def generation_status(notebook_id: str):
    status = _gen_status.get(notebook_id)
    if status and status.startswith("__slides_ready__"):
        return {"status": None, "slides_ready": True}
    return {"status": status, "slides_ready": False}


@router.get("/notebooks/{notebook_id}/artifacts/slides/history")
async def slides_history(notebook_id: str):
    history = _load_history()
    records = [r for r in history.get(notebook_id, []) if (_DOWNLOADS / r["filename"]).exists()]

    # Auto-discover merged PPTX files not yet in the database
    known = {r["filename"] for r in records}
    for path in sorted(_DOWNLOADS.glob(f"{notebook_id}_*merged*.pptx"), key=lambda p: p.stat().st_mtime):
        if path.name not in known:
            records.append({
                "filename": path.name,
                "topic": "Без темы",
                "parts": "—",
                "slides": "—",
                "created_at": datetime.fromtimestamp(path.stat().st_mtime).strftime("%d.%m.%Y %H:%M"),
            })
            known.add(path.name)

    return {"records": records}


@router.get("/notebooks/{notebook_id}/artifacts/slides/file/{filename}")
async def download_slide_file(notebook_id: str, filename: str):
    if ".." in filename or "/" in filename:
        raise HTTPException(status_code=403, detail="Access denied")
    history = _load_history()
    allowed = {r["filename"] for r in history.get(notebook_id, [])}
    # Also allow auto-discovered files that belong to this notebook
    is_owned = filename.startswith(notebook_id + "_") and filename.endswith(".pptx")
    if filename not in allowed and not is_owned:
        raise HTTPException(status_code=403, detail="Access denied")
    path = _DOWNLOADS / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        path=str(path),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=filename,
    )


@router.get("/notebooks/{notebook_id}/artifacts/audio/download")
async def download_audio(
    notebook_id: str,
    client: NotebookLMClient = Depends(get_client),
):
    out_path = _DOWNLOADS / f"{notebook_id}_podcast.mp3"
    try:
        path = await client.artifacts.download_audio(notebook_id, out_path)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))
    return FileResponse(path=str(path), media_type="audio/mpeg", filename="podcast.mp3")


@router.post("/notebooks/{notebook_id}/artifacts/{artifact_id}/delete")
async def delete_artifact(
    notebook_id: str,
    artifact_id: str,
    client: NotebookLMClient = Depends(get_client),
):
    await client.artifacts.delete(notebook_id, artifact_id)
    return RedirectResponse(url=f"/notebooks/{notebook_id}", status_code=303)
